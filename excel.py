import os
import tempfile
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


OVERVIEW = ["序号", "文件名", "相对路径", "完整路径", "所在文件夹", "文件大小", "视频时长", "分辨率", "帧率",
            "重要人物", "人物身份", "人物出现时间", "主要事件", "事件时间", "视频内容摘要", "活动/会议名称", "重要机构",
            "重要地点", "画面关键文字", "语音关键词", "检索关键词", "总体置信度", "是否需要人工复核", "处理状态", "备注"]
PEOPLE = ["人物姓名", "人物身份", "视频文件", "相对路径", "完整路径", "首次出现时间", "主要出现时间", "相关事件", "识别依据", "置信度"]
EVENTS = ["视频文件", "相对路径", "完整路径", "事件编号", "事件类型", "事件名称", "涉及人物", "开始时间", "结束时间", "事件描述", "识别依据", "置信度"]
QUALITY = ["文件名", "ffprobe状态", "场景检测状态", "关键帧状态", "OCR状态", "语音识别状态", "人物分析状态", "事件分析状态",
           "整体状态", "是否需要人工复核", "异常说明"]


def _join(values):
    return "；".join(str(v) for v in values if v not in (None, ""))


def export(records, path: Path):
    wb = Workbook()
    ws = wb.active; ws.title = "素材总览"
    people = wb.create_sheet("人物索引"); events = wb.create_sheet("事件索引"); quality = wb.create_sheet("处理质量")
    for sheet, headers in [(ws, OVERVIEW), (people, PEOPLE), (events, EVENTS), (quality, QUALITY)]:
        sheet.append(headers); sheet.freeze_panes = "A2"; sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).column_letter}1"
        sheet.sheet_view.showGridLines = False
    for i, r in enumerate(records, 1):
        f, q = r.get("file", {}), r.get("quality", {})
        ps, es = r.get("persons", []), r.get("events", [])
        notes = _join([r.get("error"), q.get("errors")])
        ws.append([i, f.get("filename"), f.get("relative_path"), f.get("path"), f.get("folder"), f.get("size"),
                   (f.get("duration") or 0) / 86400, f.get("resolution"), f.get("fps"), _join([p.get("name") for p in ps]) or "无重要人物",
                   _join([p.get("role") for p in ps]), _join([_join(p.get("time_ranges", [])) for p in ps]),
                   _join([e.get("type") for e in es]) or "无明确重要事件", _join([f'{e.get("start", "")}-{e.get("end", "")}' for e in es]),
                   r.get("summary"), _join(r.get("activities", [])), _join(r.get("organizations", [])), _join(r.get("locations", [])),
                   _join(r.get("ocr_keywords", [])), _join(r.get("speech_keywords", [])), _join(r.get("search_keywords", [])),
                   r.get("confidence"), "是" if r.get("needs_review") else "否", r.get("status"), notes])
        ws.cell(ws.max_row, 4).hyperlink = "file://" + quote(f.get("path", ""), safe="/")
        for p in ps:
            people.append([p.get("name"), p.get("role"), f.get("filename"), f.get("relative_path"), f.get("path"), p.get("first_seen"),
                           _join(p.get("time_ranges", [])), _join(p.get("related_events", [])), _join(p.get("evidence", [])), p.get("confidence")])
            people.cell(people.max_row, 5).hyperlink = "file://" + quote(f.get("path", ""), safe="/")
        for n, e in enumerate(es, 1):
            events.append([f.get("filename"), f.get("relative_path"), f.get("path"), n, e.get("type"), e.get("name"), _join(e.get("persons", [])),
                           e.get("start"), e.get("end"), e.get("description"), _join(e.get("evidence", [])), e.get("confidence")])
            events.cell(events.max_row, 3).hyperlink = "file://" + quote(f.get("path", ""), safe="/")
        quality.append([f.get("filename"), q.get("ffprobe"), q.get("scenes"), q.get("frames"), q.get("ocr"), q.get("transcription"),
                        q.get("persons"), q.get("events"), r.get("status"), "是" if r.get("needs_review") else "否", notes])
    widths = {"素材总览": [7, 24, 42, 55, 26, 14, 12, 13, 9, 24, 24, 24, 24, 24, 42, 22, 24, 20, 45, 40, 40, 12, 16, 12, 35],
              "人物索引": [22, 25, 24, 42, 55, 14, 25, 25, 45, 10], "事件索引": [24, 42, 55, 10, 16, 26, 24, 14, 14, 50, 50, 10],
              "处理质量": [24, 14, 16, 14, 14, 16, 16, 16, 14, 18, 45]}
    for sheet in wb.worksheets:
        max_row, max_col = sheet.max_row, sheet.max_column
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col, width in enumerate(widths[sheet.title], 1): sheet.column_dimensions[sheet.cell(1, col).column_letter].width = width
        sheet.auto_filter.ref = f"A1:{sheet.cell(max_row, max_col).column_letter}{max_row}"
    for cell in ws["G"][1:]: cell.number_format = "[h]:mm:ss"
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=path.parent); os.close(fd)
    try:
        wb.save(tmp); check = load_workbook(tmp, read_only=False)
        assert check.sheetnames == ["素材总览", "人物索引", "事件索引", "处理质量"]
        assert check["素材总览"].max_row == len(records) + 1
        check.close(); os.replace(tmp, path)
    finally:
        if os.path.exists(tmp): os.unlink(tmp)
