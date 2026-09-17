"""Read the eight manually revised beta workbooks into a neutral JSON payload.

This script is deliberately read-only with respect to the source workbooks.  It
does not open them for writing or save them again.
"""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


SOURCE_DIR = Path("/Volumes/公用文件夹/WDCC2026/视频文件检索目录beta版")
OUT = Path("/tmp/wdcc_beta_merge/beta_merge_data.json")

OVERVIEW_HEADERS = [
    "序号", "来源表格", "来源Sheet", "所属批次 / 文件夹", "文件名", "相对路径", "完整路径", "所在文件夹",
    "文件大小", "视频时长", "分辨率", "帧率", "重要人物", "人物身份", "人物出现时间", "主要事件",
    "事件发生时间", "视频内容摘要", "活动/会议名称", "重要机构", "重要地点", "重要日期", "画面关键文字",
    "语音关键词", "检索关键词", "总体置信度", "是否需要人工复核", "处理状态", "备注",
]
PERSON_HEADERS = [
    "人物姓名/描述", "人物身份", "视频文件", "所属批次 / 文件夹", "相对路径", "完整路径", "首次出现时间",
    "主要出现时间", "相关事件", "识别依据", "置信度", "来源表格", "来源Sheet",
]
EVENT_HEADERS = [
    "视频文件", "所属批次 / 文件夹", "相对路径", "完整路径", "事件编号", "事件类型", "事件名称", "涉及人物",
    "开始时间", "结束时间", "事件描述", "识别依据", "置信度", "来源表格", "来源Sheet",
]
QUALITY_HEADERS = [
    "文件名", "所属批次 / 文件夹", "完整路径", "ffprobe状态", "场景检测状态", "关键帧状态", "OCR状态",
    "语音识别状态", "人物识别状态", "事件识别状态", "整体状态", "是否需要人工复核", "异常说明", "来源表格", "来源Sheet",
]


def clean(value):
    if value is None:
        return ""
    return value


def rows_from_sheet(ws):
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return [], []
    headers = [str(clean(v)).strip() for v in values[0]]
    rows = []
    for values_row in values[1:]:
        row = {headers[i]: clean(values_row[i]) if i < len(values_row) else "" for i in range(len(headers))}
        if any(v not in (None, "") for v in row.values()):
            rows.append(row)
    return headers, rows


def value(row, *names):
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return ""


def batch_name(file_name):
    mapping = {
        "视频素材检索结果_蔡坤测试.xlsx": "蔡坤",
        "视频素材检索结果_顾珣.xlsx": "顾珣",
        "视频素材检索结果_邱语.xlsx": "邱语",
        "视频素材检索结果_视频总结片_9.19-9.25.xlsx": "视频总结片_9.19-9.25",
        "视频素材检索结果_视频总结片_9.27.xlsx": "视频总结片_9.27",
        "视频素材检索结果_视频总结片_9.28.xlsx": "视频总结片_9.28",
        "视频素材检索结果_视频总结片_剩余2文件夹.xlsx": "视频总结片_剩余2文件夹",
        "WDCC三部长视频重大事件索引.xlsx": "WDCC三部长视频",
    }
    return mapping[file_name]


def overview_record(row, source, sheet, batch):
    return {
        "来源表格": source, "来源Sheet": sheet,
        "所属批次 / 文件夹": value(row, "所属子文件夹") or batch,
        "文件名": value(row, "文件名", "视频名称"),
        "相对路径": value(row, "相对路径"), "完整路径": value(row, "完整路径"),
        "所在文件夹": value(row, "所在文件夹"), "文件大小": value(row, "文件大小"),
        "视频时长": value(row, "视频时长"), "分辨率": value(row, "分辨率"), "帧率": value(row, "帧率"),
        "重要人物": value(row, "重要人物"), "人物身份": value(row, "人物身份"),
        "人物出现时间": value(row, "人物出现时间"), "主要事件": value(row, "主要事件"),
        "事件发生时间": value(row, "事件时间", "事件发生时间"), "视频内容摘要": value(row, "视频内容摘要"),
        "活动/会议名称": value(row, "活动/会议名称"), "重要机构": value(row, "重要机构"),
        "重要地点": value(row, "重要地点"), "重要日期": value(row, "重要日期"),
        "画面关键文字": value(row, "画面关键文字"), "语音关键词": value(row, "语音关键词"),
        "检索关键词": value(row, "检索关键词"), "总体置信度": value(row, "总体置信度"),
        "是否需要人工复核": value(row, "是否需要人工复核"), "处理状态": value(row, "处理状态", "整体状态"),
        "备注": value(row, "备注"),
    }


def person_record(row, source, sheet, batch):
    return {
        "人物姓名/描述": value(row, "人物姓名/描述", "人物姓名"), "人物身份": value(row, "人物身份"),
        "视频文件": value(row, "视频文件", "视频名称"), "所属批次 / 文件夹": value(row, "所属子文件夹") or batch,
        "相对路径": value(row, "相对路径"), "完整路径": value(row, "完整路径"),
        "首次出现时间": value(row, "首次出现时间"), "主要出现时间": value(row, "主要出现时间", "主要出现时间段"),
        "相关事件": value(row, "相关事件"), "识别依据": value(row, "识别依据"), "置信度": value(row, "置信度"),
        "来源表格": source, "来源Sheet": sheet,
    }


def event_record(row, source, sheet, batch):
    return {
        "视频文件": value(row, "视频文件", "视频名称"), "所属批次 / 文件夹": value(row, "所属子文件夹") or batch,
        "相对路径": value(row, "相对路径"), "完整路径": value(row, "完整路径"),
        "事件编号": value(row, "事件编号"), "事件类型": value(row, "事件类型"),
        "事件名称": value(row, "事件名称", "事件标题"), "涉及人物": value(row, "涉及人物", "重要人物"),
        "开始时间": value(row, "开始时间"), "结束时间": value(row, "结束时间"),
        "事件描述": value(row, "事件描述"), "识别依据": value(row, "识别依据"), "置信度": value(row, "置信度"),
        "来源表格": source, "来源Sheet": sheet,
    }


def quality_record(row, source, sheet, batch):
    return {
        "文件名": value(row, "文件名", "视频名称"), "所属批次 / 文件夹": value(row, "所属子文件夹") or batch,
        "完整路径": value(row, "完整路径"), "ffprobe状态": value(row, "ffprobe状态"),
        "场景检测状态": value(row, "场景检测状态"), "关键帧状态": value(row, "关键帧状态", "关键帧数量"),
        "OCR状态": value(row, "OCR状态"), "语音识别状态": value(row, "语音识别状态", "ASR状态"),
        "人物识别状态": value(row, "人物识别状态"), "事件识别状态": value(row, "事件识别状态", "事件分区状态"),
        "整体状态": value(row, "整体状态"), "是否需要人工复核": value(row, "是否需要人工复核", "是否存在需要人工复核的事件"),
        "异常说明": value(row, "异常说明"), "来源表格": source, "来源Sheet": sheet,
    }


def cell_diff(a, b):
    keys = [k for k in a if k not in {"来源表格", "来源Sheet", "序号"}]
    return [k for k in keys if str(a.get(k, "")) != str(b.get(k, ""))]


def main():
    files = sorted(SOURCE_DIR.glob("*.xlsx"))
    if len(files) != 8:
        raise RuntimeError(f"Expected exactly 8 .xlsx source files; found {len(files)}")
    payload = {"overview_headers": OVERVIEW_HEADERS, "person_headers": PERSON_HEADERS, "event_headers": EVENT_HEADERS,
               "quality_headers": QUALITY_HEADERS, "overview": [], "persons": [], "events": [], "quality": [],
               "sources": [], "duplicates": []}
    for path in files:
        source = path.name
        batch = batch_name(source)
        wb = load_workbook(path, read_only=True, data_only=False)
        read_counts, import_counts, skip_counts, notes = [], [], [], []
        long_quality_rows = []
        for ws in wb.worksheets:
            headers, rows = rows_from_sheet(ws)
            read_counts.append(f"{ws.title}:{len(rows)}")
            imported = 0
            skipped = 0
            if ws.title == "素材总览":
                for row in rows:
                    payload["overview"].append(overview_record(row, source, ws.title, batch)); imported += 1
            elif ws.title in {"人物索引", "重要人物索引"}:
                for row in rows:
                    payload["persons"].append(person_record(row, source, ws.title, batch)); imported += 1
            elif ws.title in {"事件索引", "关键事件明细", "重大事件总时间轴"}:
                for row in rows:
                    payload["events"].append(event_record(row, source, ws.title, batch)); imported += 1
            elif ws.title == "处理质量":
                for row in rows:
                    rec = quality_record(row, source, ws.title, batch)
                    payload["quality"].append(rec); imported += 1
                    if source == "WDCC三部长视频重大事件索引.xlsx":
                        long_quality_rows.append(row)
            elif source == "WDCC三部长视频重大事件索引.xlsx" and ws.title.startswith("#"):
                skipped = len(rows)
                notes.append(f"{ws.title}与“重大事件总时间轴”内容重复，未重复导入")
            else:
                skipped = len(rows)
                notes.append(f"{ws.title}未映射至标准工作表")
            import_counts.append(f"{ws.title}:{imported}")
            if skipped:
                skip_counts.append(f"{ws.title}:{skipped}")
        # The long-video source has no material overview.  Its three manually revised
        # quality rows are the only available non-event source for those videos.
        if source == "WDCC三部长视频重大事件索引.xlsx":
            for row in long_quality_rows:
                payload["overview"].append(overview_record(row, source, "处理质量", batch))
            notes.append("该源表无“素材总览”；已由其处理质量中的3条视频记录建立最小素材总览行，未生成任何新内容")
        payload["sources"].append({
            "source": source, "path": str(path), "sheets": "；".join(ws.title for ws in wb.worksheets),
            "read_rows": "；".join(read_counts), "imported_rows": "；".join(import_counts),
            "skipped_rows": "；".join(skip_counts), "notes": "；".join(notes),
        })

    groups = defaultdict(list)
    for rec in payload["overview"]:
        path = str(rec["完整路径"]).strip()
        key = ("path", path) if path else ("fallback", str(rec["相对路径"]).strip(), str(rec["文件名"]).strip(), str(rec["视频时长"]).strip())
        groups[key].append(rec)
    for _, records in groups.items():
        if len(records) < 2:
            continue
        for idx in range(1, len(records)):
            base, current = records[0], records[idx]
            diffs = cell_diff(base, current)
            payload["duplicates"].append({
                "文件名": current["文件名"], "路径A": base["完整路径"] or base["相对路径"], "来源表A": base["来源表格"], "来源SheetA": base["来源Sheet"],
                "路径B": current["完整路径"] or current["相对路径"], "来源表B": current["来源表格"], "来源SheetB": current["来源Sheet"],
                "差异字段": "；".join(diffs), "是否完全一致": "是" if not diffs else "否",
                "建议人工检查": "否（完全一致）" if not diffs else "是（保留所有来源记录，未自动删除）",
            })
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, default=str), encoding="utf-8")
    print(json.dumps({k: len(payload[k]) for k in ("overview", "persons", "events", "quality", "sources", "duplicates")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
