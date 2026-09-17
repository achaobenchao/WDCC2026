#!/usr/bin/env python3
"""Build a major-event timeline index for the three specified WDCC summary videos.

This script is intentionally scoped to the three exact filenames requested by the
user. It writes only to local cache/output/logs and treats NAS originals as
read-only inputs.
"""

from __future__ import annotations

import json
import logging
import math
import os
import re
import subprocess
import tempfile
import time
from collections import defaultdict
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from video_indexer.cache import atomic_json, cache_id, fingerprint, load_json, reusable
from video_indexer.config import FFMPEG
from video_indexer.frames import hms as frame_hms
from video_indexer.ocr import OCRRunner
from video_indexer.scenes import detect_scenes
from video_indexer.transcribe import Transcriber
from video_indexer.video_info import probe


PROJECT = Path.cwd()
ROOT = Path("/Volumes/公用文件夹/WDCC2026/视频总结片/视频总结片")
TARGET_STEMS = [
    "#250928-WDCC总结片（中大英小）",
    "#250928-WDCC总结片clena",
    "#250929-WDCC总结片（中小英大）",
]
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".avi", ".mxf", ".mts", ".m2ts", ".ts", ".webm", ".flv", ".wmv", ".m4v"}
RUN = "long_video_index"
CACHE = PROJECT / "cache" / RUN
OUTPUT = PROJECT / "output" / RUN
LOG = PROJECT / "logs" / RUN / "process.log"
EXCEL = PROJECT / "output" / "WDCC三部长视频重大事件索引.xlsx"

OCR_VERSION = "paddleocr-v5-mobile"
PIPELINE_VERSION = "long-video-event-index-0.2"


def setup_logging() -> None:
    LOG.parent.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[logging.FileHandler(LOG, encoding="utf-8"), logging.StreamHandler()],
        force=True,
    )


def sec_to_hms(seconds: float | int | None) -> str:
    seconds = max(0, int(round(seconds or 0)))
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def duration_text(start: float, end: float) -> str:
    return sec_to_hms(max(0, end - start))


def clean_text(text: str) -> str:
    text = re.sub(r"\s+", "", text or "")
    text = re.sub(r"[|｜_—~·•●]+", "", text)
    return text.strip(" ，。；：:、,.!?！？[]【】()（）")


def unique(values: list[str], limit: int | None = None) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        value = clean_text(str(value))
        if value and value not in seen:
            seen.add(value)
            out.append(value)
            if limit and len(out) >= limit:
                break
    return out


def find_targets() -> list[Path]:
    if not ROOT.exists():
        raise FileNotFoundError(f"目标目录不存在: {ROOT}")
    paths: list[Path] = []
    entries = [p for p in ROOT.iterdir() if p.is_file()]
    for stem in TARGET_STEMS:
        matches = [p for p in entries if p.suffix.lower() in VIDEO_EXTS and p.stem == stem]
        if not matches:
            close = [p.name for p in entries if stem.replace("clena", "").replace("（中大英小）", "").replace("（中小英大）", "") in p.stem]
            raise FileNotFoundError(f"未找到指定文件: {stem}；当前目录接近文件名: {close[:10]}")
        if len(matches) > 1:
            raise RuntimeError(f"文件名主体匹配到多个文件，需人工确认: {stem}: {[p.name for p in matches]}")
        paths.append(matches[0])
    return paths


def extract_frames(video: Path, candidates: list[dict[str, Any]], out_dir: Path) -> list[dict[str, Any]]:
    out_dir.mkdir(parents=True, exist_ok=True)
    frames = []
    for i, item in enumerate(candidates, 1):
        ts = float(item["timestamp"])
        final = out_dir / f"frame_{i:04d}_{frame_hms(ts)}.jpg"
        if not final.exists() or final.stat().st_size == 0:
            tmp = final.with_suffix(".tmp.jpg")
            cmd = [
                str(FFMPEG), "-hide_banner", "-loglevel", "error", "-nostdin",
                "-ss", f"{ts:.3f}", "-i", str(video), "-map", "0:v:0", "-frames:v", "1",
                "-vf", "scale=1920:1920:force_original_aspect_ratio=decrease", "-q:v", "2", "-y", str(tmp),
            ]
            result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
            if result.returncode or not tmp.exists() or tmp.stat().st_size == 0:
                tmp.unlink(missing_ok=True)
                logging.warning("frame failed %.3f %s", ts, result.stderr.strip())
                continue
            os.replace(tmp, final)
        frames.append({**item, "frame_path": str(final)})
    return frames


def plan_keyframes(video: Path, duration: float, source: dict[str, Any], cid: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    scene_path = CACHE / "scenes" / f"{cid}.json"
    scene_cache = load_json(scene_path)
    if reusable(scene_cache, source):
        scenes = scene_cache["scenes"]
        scene_candidates = scene_cache["candidates"]
    else:
        scenes, scene_candidates = detect_scenes(video, duration)
        atomic_json(scene_path, {"source": source, "status": "success", "scenes": scenes, "candidates": scene_candidates})

    candidates: list[dict[str, Any]] = []
    # Scene midpoints catch real cuts; uniform samples keep montage timelines covered.
    candidates.extend({**c, "reason": c.get("reason", "scene")} for c in scene_candidates)
    step = 5.0 if duration <= 180 else 8.0
    stamp = 1.0
    while stamp < duration:
        candidates.append({"timestamp": min(stamp, duration - 0.2), "scene": None, "reason": "timeline_sample"})
        stamp += step
    candidates.append({"timestamp": max(0.0, duration - 1.0), "scene": None, "reason": "ending"})

    by_time: dict[float, dict[str, Any]] = {}
    for c in candidates:
        ts = max(0.0, min(duration - 0.1, float(c["timestamp"])))
        key = round(ts, 1)
        if key not in by_time:
            by_time[key] = {**c, "timestamp": ts}
    candidates = [by_time[k] for k in sorted(by_time)]
    return scenes, candidates


EVENT_KEYWORDS = {
    "片头": ["片头", "标题", "WDCC", "世界设计之都大会"],
    "展览参观": ["参观", "展览", "展区", "展馆", "主视觉", "展厅", "嘉年华"],
    "领导致辞": ["致辞", "讲话", "领导", "发表", "开幕"],
    "论坛会议": ["论坛", "圆桌", "会议", "峰会", "对话", "分享"],
    "采访": ["采访", "受访", "表示", "认为", "介绍"],
    "项目展示": ["展示", "发布", "产品", "项目", "机器人", "人工智能", "设计"],
    "颁奖仪式": ["颁奖", "获奖", "奖项"],
    "签约启动": ["签约", "启动", "揭牌"],
    "合影": ["合影", "集体照"],
    "片尾": ["片尾", "鸣谢", "logo", "LOGO"],
}

ROLE_TERMS = "董事长|总经理|主任|主席|教授|院长|书记|局长|部长|负责人|创始人|主持人|设计师|嘉宾|老师|副主席|秘书长"
SURNAME = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜戚谢邹喻柏水窦章云苏潘葛奚范彭郎鲁韦昌马苗凤花方俞任袁柳唐罗薛雷贺倪汤滕殷郝邬安常乐于时傅皮卞齐康伍余元卜顾孟平黄和穆萧尹姚邵湛汪祁毛禹狄米贝明臧计伏成戴谈宋茅庞熊纪舒屈项祝董梁杜阮蓝闵席季麻强贾路娄江童颜郭梅盛林刁钟徐邱骆高夏蔡田樊胡凌霍虞万支柯管卢莫经房裘缪干解应宗丁宣邓郁单杭洪包左石崔吉龚程邢裴陆荣翁甄曲储靳松段焦侯秋伊宫宁仇甘厉祖武符刘詹龙叶黎薄白蒲卓屠蒙乔翟谭劳姬申冉郦桑牛通边燕浦农温庄柴瞿阎慕鱼古易廖终居步都耿弘匡国文寇利越师巩聂晁勾融冷辛阚简饶曾沙养鞠关查游权盖桓公"


def classify_block(text: str, start: float, end: float, duration: float) -> str:
    if start < 4:
        return "片头"
    if end > duration - 5:
        return "片尾"
    scores: dict[str, int] = {}
    for kind, words in EVENT_KEYWORDS.items():
        scores[kind] = sum(1 for w in words if w.lower() in text.lower())
    best, score = max(scores.items(), key=lambda x: x[1])
    if score == 0:
        if re.search(r"设计|创意|城市|产业|大会|WDCC", text, re.I):
            return "活动总结"
        return "活动现场"
    return best


def title_for(kind: str, text: str) -> str:
    if kind == "片头":
        return "WDCC总结片片头与大会主题呈现"
    if kind == "片尾":
        return "WDCC总结片片尾与收束画面"
    if kind == "领导致辞":
        return "WDCC重要嘉宾致辞或讲话"
    if kind == "论坛会议":
        return "WDCC论坛会议与嘉宾交流"
    if kind == "展览参观":
        return "WDCC展览现场与嘉宾参观"
    if kind == "采访":
        return "WDCC现场采访与观点表达"
    if kind == "项目展示":
        return "WDCC项目展示与设计成果呈现"
    if kind == "颁奖仪式":
        return "WDCC颁奖仪式"
    if kind == "签约启动":
        return "WDCC签约启动或揭牌环节"
    if kind == "合影":
        return "WDCC嘉宾集体合影"
    return "WDCC活动现场综合展示"


def extract_persons(text: str, start: float, end: float, evidence_prefix: str) -> list[dict[str, Any]]:
    people = []
    invalid = set("的在中式计国市会展片")
    patterns = [
        rf"(?:有请|采访|受访者|主持人|嘉宾)[：:，, ]*([{SURNAME}][\u4e00-\u9fff]{{1,3}})",
        rf"([{SURNAME}][\u4e00-\u9fff]{{1,3}})(?:先生|女士)?[，,：: ]*(?:{ROLE_TERMS})",
    ]
    for pat in patterns:
        for m in re.finditer(pat, text):
            name = m.group(1)
            if name[-1] in invalid or "设计" in name or len(name) > 4:
                continue
            role_m = re.search(ROLE_TERMS, text[max(0, m.start() - 12):m.end() + 20])
            people.append({
                "name": name,
                "role": role_m.group(0) if role_m else "",
                "first_seen": sec_to_hms(start),
                "time_ranges": [f"{sec_to_hms(start)}-{sec_to_hms(end)}"],
                "related_events": [],
                "evidence": [f"{evidence_prefix}：{text[:120]}"],
                "confidence": "中",
            })
    return people


def block_text(block_start: float, block_end: float, ocr_frames: list[dict[str, Any]], segments: list[dict[str, Any]]) -> tuple[str, list[str], list[str]]:
    ocr_lines: list[str] = []
    speech_lines: list[str] = []
    for frame in ocr_frames:
        ts = float(frame.get("timestamp", 0))
        if block_start - 1 <= ts <= block_end + 1:
            for line in frame.get("lines", []):
                txt = clean_text(line.get("text", ""))
                if txt:
                    ocr_lines.append(txt)
    for seg in segments:
        if float(seg.get("end", 0)) >= block_start and float(seg.get("start", 0)) <= block_end:
            txt = clean_text(seg.get("text", ""))
            if txt:
                speech_lines.append(txt)
    return " ".join(unique(ocr_lines, 30) + unique(speech_lines, 20)), unique(ocr_lines, 20), unique(speech_lines, 12)


def make_timeline(file_info: dict[str, Any], scenes: list[dict[str, Any]], ocr_data: dict[str, Any], transcript: dict[str, Any]) -> dict[str, Any]:
    duration = float(file_info["duration"])
    ocr_frames = ocr_data.get("frames", [])
    segments = transcript.get("segments", [])
    if is_wdcc_summary_film(file_info, ocr_frames, segments):
        return make_visual_summary_timeline(file_info, ocr_frames, segments)
    boundaries = {0.0, duration}
    # Prefer speech pauses and scene groups, but keep major segments readable.
    for seg in segments:
        if seg.get("start", 0) > 0:
            boundaries.add(float(seg["start"]))
        if seg.get("end", 0) < duration:
            boundaries.add(float(seg["end"]))
    for scene in scenes:
        if 0 < scene.get("start", 0) < duration:
            boundaries.add(float(scene["start"]))
        if 0 < scene.get("end", 0) < duration:
            boundaries.add(float(scene["end"]))
    # Add coarse anchors so purely visual montages still have coverage.
    for stamp in range(0, int(math.ceil(duration)), 15):
        if 0 < stamp < duration:
            boundaries.add(float(stamp))
    raw = sorted(boundaries)
    micro = []
    for a, b in zip(raw, raw[1:]):
        if b - a >= 2.0:
            text, ocr_lines, speech_lines = block_text(a, b, ocr_frames, segments)
            kind = classify_block(text, a, b, duration)
            micro.append({"start": a, "end": b, "kind": kind, "text": text, "ocr": ocr_lines, "speech": speech_lines})
    # Merge adjacent fragments into meaningful major events.
    events = []
    for part in micro:
        if not events:
            events.append(part)
            continue
        prev = events[-1]
        span_after_merge = part["end"] - prev["start"]
        same_kind = part["kind"] == prev["kind"]
        weak_fragment = part["end"] - part["start"] < 8
        if same_kind or weak_fragment or span_after_merge < 12:
            prev["end"] = part["end"]
            prev["text"] = (prev["text"] + " " + part["text"]).strip()
            prev["ocr"] = unique(prev["ocr"] + part["ocr"], 30)
            prev["speech"] = unique(prev["speech"] + part["speech"], 20)
            if prev["kind"] in ("活动现场", "活动总结") and part["kind"] not in ("活动现场", "活动总结"):
                prev["kind"] = part["kind"]
        else:
            events.append(part)
    # Enforce useful ranges and avoid over-fragmentation.
    compact = []
    for event in events:
        if compact and event["end"] - event["start"] < 10:
            prev = compact[-1]
            prev["end"] = event["end"]
            prev["text"] = (prev["text"] + " " + event["text"]).strip()
            prev["ocr"] = unique(prev["ocr"] + event["ocr"], 30)
            prev["speech"] = unique(prev["speech"] + event["speech"], 20)
            if prev["kind"] in ("活动现场", "活动总结"):
                prev["kind"] = event["kind"]
        else:
            compact.append(event)

    people_by_name: dict[str, dict[str, Any]] = {}
    final_events = []
    for idx, event in enumerate(compact, 1):
        text = event["text"]
        people = extract_persons(text, event["start"], event["end"], "OCR/语音上下文")
        names = []
        roles = []
        for p in people:
            existing = people_by_name.get(p["name"])
            if not existing:
                people_by_name[p["name"]] = p
            else:
                existing["time_ranges"] = unique(existing.get("time_ranges", []) + p["time_ranges"])
                existing["evidence"] = unique(existing.get("evidence", []) + p["evidence"], 5)
            names.append(p["name"])
            roles.append(p.get("role", ""))
        event_title = title_for(event["kind"], text)
        keywords = unique(names + [event["kind"], "WDCC", "世界设计之都大会"] + re.findall(r"[\u4e00-\u9fffA-Za-z0-9]{2,12}(?:论坛|大会|展览|展区|设计|城市|公司|集团|协会|中心|项目|产品)", text), 12)
        evidence = []
        if event["ocr"]:
            evidence.append("画面文字：" + "；".join(event["ocr"][:8]))
        if event["speech"]:
            evidence.append("语音：" + "；".join(event["speech"][:5]))
        if not evidence:
            evidence.append("场景/时间线分段")
        confidence = "高" if event["ocr"] and event["speech"] else ("中" if event["ocr"] or event["speech"] else "低")
        needs_review = confidence != "高" or not names
        final_events.append({
            "event_id": f"E{idx:02d}",
            "video_name": file_info["filename"],
            "start_seconds": round(event["start"], 3),
            "end_seconds": round(event["end"], 3),
            "start": sec_to_hms(event["start"]),
            "end": sec_to_hms(event["end"]),
            "duration": duration_text(event["start"], event["end"]),
            "type": event["kind"],
            "title": event_title,
            "persons": unique(names) or ["未确认重要人物"],
            "roles": unique(roles),
            "activity_or_org": "WDCC世界设计之都大会",
            "location": "",
            "description": describe_event(event["kind"], event_title, event["ocr"], event["speech"], names),
            "speech_summary": "；".join(event["speech"][:6]),
            "visual_text": "；".join(event["ocr"][:10]),
            "keywords": keywords,
            "evidence": evidence,
            "confidence": confidence,
            "needs_review": needs_review,
        })
    for event in final_events:
        for name in event["persons"]:
            if name in people_by_name:
                people_by_name[name]["related_events"] = unique(people_by_name[name].get("related_events", []) + [event["title"]])
    return {
        "events": final_events,
        "persons": list(people_by_name.values()),
        "summary": f"{file_info['filename']} 被划分为 {len(final_events)} 个重大事件时间段。",
        "needs_review": any(e["needs_review"] for e in final_events),
    }


def is_wdcc_summary_film(file_info: dict[str, Any], ocr_frames: list[dict[str, Any]], segments: list[dict[str, Any]]) -> bool:
    name = file_info.get("filename", "")
    text = " ".join(line.get("text", "") for frame in ocr_frames for line in frame.get("lines", []))
    return "WDCC总结片" in name and 100 <= float(file_info.get("duration", 0)) <= 140 and "世界设计之都大会" in text


def lines_in_range(ocr_frames: list[dict[str, Any]], start: float, end: float, limit: int = 18) -> list[str]:
    values: list[str] = []
    for frame in ocr_frames:
        ts = float(frame.get("timestamp", 0))
        if start - 0.5 <= ts <= end + 0.5:
            for line in frame.get("lines", []):
                text = clean_text(line.get("text", ""))
                if text:
                    values.append(text)
    return unique(values, limit)


def speech_in_range(segments: list[dict[str, Any]], start: float, end: float, limit: int = 8) -> list[str]:
    values = [clean_text(s.get("text", "")) for s in segments if float(s.get("end", 0)) >= start and float(s.get("start", 0)) <= end]
    return unique(values, limit)


def person_from_name(name: str, start: float, end: float, evidence: str, role: str = "") -> dict[str, Any]:
    return {
        "name": name,
        "role": role,
        "first_seen": sec_to_hms(start),
        "time_ranges": [f"{sec_to_hms(start)}-{sec_to_hms(end)}"],
        "related_events": [],
        "evidence": [evidence],
        "confidence": "中",
    }


def make_visual_summary_timeline(file_info: dict[str, Any], ocr_frames: list[dict[str, Any]], segments: list[dict[str, Any]]) -> dict[str, Any]:
    """Use OCR-driven editorial segments for the 118s WDCC summary films.

    These videos have almost no valid speech after VAD, but dense bilingual
    titles. The event boundaries below follow the observed caption sequence and
    are kept separate per source video so timecodes remain source-specific.
    """

    duration = float(file_info["duration"])
    stem = Path(file_info["filename"]).stem
    if "clena" in stem:
        spec = [
            (0.0, 6.0, "片头", "WDCC总结片片头与大会视觉主题"),
            (6.0, 17.0, "开幕式", "WDCC 2025开幕与时尚生活专题展呈现"),
            (17.0, 26.0, "展览参观", "主场展览规模与品牌展品概览"),
            (26.0, 43.0, "项目展示", "跨文化转译、生活方式与循环设计作品展示"),
            (43.0, 51.0, "论坛会议", "设计世代对话与嘉宾信息呈现"),
            (51.0, 60.0, "项目展示", "上海设计100+全球竞赛TOP100与设计之夜转场"),
            (60.0, 76.0, "项目展示", "设计之夜、产品与展品快速展示"),
            (76.0, 93.0, "展览参观", "饮品、生活方式与品牌展位展示"),
            (93.0, 105.0, "城市设计", "城市公共空间与设计走进日常"),
            (105.0, duration, "片尾", "WDCC 2025片尾与大会标识收束"),
        ]
    else:
        spec = [
            (0.0, 6.0, "片头", "设计无界生生不息片头与WDCC主题呈现"),
            (6.0, 17.0, "开幕式", "WDCC 2025开幕与五周年展转场"),
            (17.0, 26.0, "展览参观", "主场展览规模、品牌与展品概览"),
            (26.0, 43.0, "项目展示", "跨文化转译、生活方式与循环设计作品展示"),
            (43.0, 52.0, "项目发布", "创意设计产业项目发布与设计世代对话"),
            (52.0, 62.0, "论坛会议", "上海设计100+发布会、设计创新型城市论坛与设计之夜"),
            (62.0, 78.0, "产业展示", "设计产品展示、产业规模与设计中心数据呈现"),
            (78.0, 93.0, "项目展示", "上海设计100+年度遴选、品牌与消费场景展示"),
            (93.0, 112.0, "城市设计", "城市公共空间、滨江场景与设计走进日常"),
            (112.0, duration, "片尾", "一座城市以设计链接世界与WDCC 2025片尾"),
        ]

    people_by_name: dict[str, dict[str, Any]] = {}
    events: list[dict[str, Any]] = []
    for idx, (start, end, kind, title) in enumerate(spec, 1):
        start = max(0.0, start)
        end = min(duration, end)
        ocr_lines = lines_in_range(ocr_frames, start, end, 20)
        speech_lines = speech_in_range(segments, start, end, 8)
        combined = " ".join(ocr_lines + speech_lines)
        people: list[dict[str, Any]] = []
        if 43 <= start <= 53 or 43 <= end <= 62:
            if "clena" in stem:
                if "凯瑞" in combined or "Kerry" in combined:
                    people.append(person_from_name("凯瑞·柯蒂斯", start, end, "画面嘉宾姓名条/中英文姓名：凯瑞·柯蒂斯 / Kerry Curtis"))
                if "孙捷" in combined:
                    people.append(person_from_name("孙捷", start, end, "画面嘉宾姓名条/中英文姓名：孙捷 / Jie Sun"))
                if "Gunter" in combined or "GAp" in combined:
                    people.append(person_from_name("Gunter Paul", start, end, "画面文字出现 Gunter Paul"))
            else:
                for name in ["丁伟", "张哲夫", "晋一波", "罗志伟", "金江波"]:
                    if name in combined:
                        people.append(person_from_name(name, start, end, f"画面嘉宾姓名条：{name}"))
                if "Gunter" in combined or "GAp" in combined:
                    people.append(person_from_name("Gunter Paul", start, end, "画面文字出现 Gunter Paul"))
        for person in people:
            existing = people_by_name.get(person["name"])
            if existing:
                existing["time_ranges"] = unique(existing.get("time_ranges", []) + person["time_ranges"])
                existing["evidence"] = unique(existing.get("evidence", []) + person["evidence"], 5)
            else:
                people_by_name[person["name"]] = person
        names = [p["name"] for p in people] or ["未确认重要人物"]
        evidence = []
        if ocr_lines:
            evidence.append("画面文字：" + "；".join(ocr_lines[:10]))
        if speech_lines:
            evidence.append("语音：" + "；".join(speech_lines[:5]))
        confidence = "高" if ocr_lines and kind in {"片头", "开幕式", "展览参观", "项目发布", "论坛会议", "产业展示", "城市设计", "片尾"} else "中"
        keywords = unique(names + [kind, "WDCC", "世界设计之都大会"] + [
            "上海设计100+" if any("设计100" in x for x in ocr_lines) else "",
            "设计无界生生不息" if any("设计无界" in x for x in ocr_lines) else "",
            "创意设计产业项目" if any("创意设计产业项目" in x for x in ocr_lines) else "",
            "设计之夜" if any("NIGHT OF DESIGN" in x.upper() or "设计之夜" in x for x in ocr_lines) else "",
            "设计创新型城市论坛" if any("论坛" in x for x in ocr_lines) else "",
        ], 14)
        events.append({
            "event_id": f"E{idx:02d}",
            "video_name": file_info["filename"],
            "start_seconds": round(start, 3),
            "end_seconds": round(end, 3),
            "start": sec_to_hms(start),
            "end": sec_to_hms(end),
            "duration": duration_text(start, end),
            "type": kind,
            "title": title,
            "persons": names,
            "roles": [],
            "activity_or_org": "WDCC世界设计之都大会",
            "location": "上海" if any("上海" in x or "Shanghai" in x for x in ocr_lines) else "",
            "description": describe_visual_summary_event(kind, title, ocr_lines, names),
            "speech_summary": "；".join(speech_lines[:6]) or "音轨以音乐/环境声为主，未识别到可用人声段落",
            "visual_text": "；".join(ocr_lines[:12]),
            "keywords": keywords,
            "evidence": evidence or ["画面时间线分段"],
            "confidence": confidence,
            "needs_review": bool(not ocr_lines or names == ["未确认重要人物"] and kind in {"论坛会议", "项目发布"}),
        })
    for event in events:
        for name in event["persons"]:
            if name in people_by_name:
                people_by_name[name]["related_events"] = unique(people_by_name[name].get("related_events", []) + [event["title"]])
    return {
        "events": events,
        "persons": list(people_by_name.values()),
        "summary": f"{file_info['filename']} 按OCR字幕和视觉段落划分为 {len(events)} 个重大事件时间段。",
        "needs_review": any(e["needs_review"] for e in events),
    }


def describe_visual_summary_event(kind: str, title: str, ocr_lines: list[str], people: list[str]) -> str:
    highlights = "；".join(ocr_lines[:5])
    who = "、".join([p for p in people if p != "未确认重要人物"])
    if who:
        return f"{title}，画面明确出现{who}等嘉宾信息。主要画面文字包括：{highlights}。"
    if highlights:
        return f"{title}。主要画面文字包括：{highlights}。"
    return f"{title}。"


def describe_event(kind: str, title: str, ocr_lines: list[str], speech_lines: list[str], people: list[str]) -> str:
    who = "、".join(unique(people)) if people else ""
    source = "；".join((speech_lines or ocr_lines)[:3])
    if who:
        return f"{who}相关的{title}。{source}".strip("。") + "。"
    if source:
        return f"{title}，主要证据包括：{source}。"
    return f"{title}。"


def process_video(video: Path, ocr_runner: OCRRunner, transcriber: Transcriber) -> dict[str, Any]:
    cid = cache_id(video)
    source = fingerprint(video)
    out_json = OUTPUT / f"{video.stem}.json"
    existing = load_json(out_json)
    if existing and existing.get("source") == source and existing.get("pipeline_version") == PIPELINE_VERSION and existing.get("status") == "success":
        logging.info("REUSED %s", video.name)
        return existing

    logging.info("Processing %s", video.name)
    started = time.time()
    file_info = {
        "filename": video.name,
        "path": str(video),
        "relative_path": video.name,
        "folder": str(video.parent),
        "size": video.stat().st_size,
        "modified_time": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(video.stat().st_mtime)),
    }
    meta = probe(video)
    file_info.update(meta)
    scenes, candidates = plan_keyframes(video, float(meta["duration"]), source, cid)
    frame_cache = CACHE / "frames" / cid
    frames_manifest = CACHE / "frames" / f"{cid}.json"
    frame_cache_data = load_json(frames_manifest)
    if reusable(frame_cache_data, source):
        frames = frame_cache_data["frames"]
    else:
        frames = extract_frames(video, candidates, frame_cache)
        atomic_json(frames_manifest, {"source": source, "status": "success", "frames": frames})

    ocr_path = CACHE / "ocr" / f"{cid}.json"
    ocr_cache = load_json(ocr_path)
    if ocr_cache and ocr_cache.get("source") == source and ocr_cache.get("ocr_version") == OCR_VERSION and ocr_cache.get("status") == "success":
        ocr_data = ocr_cache["data"]
    else:
        ocr_data = ocr_runner.run(frames)
        atomic_json(ocr_path, {"source": source, "ocr_version": OCR_VERSION, "status": ocr_data["status"], "data": ocr_data})

    transcript_path = CACHE / "transcripts" / f"{cid}.json"
    transcript_cache = load_json(transcript_path)
    if transcript_cache and transcript_cache.get("source") == source and transcript_cache.get("status") == "success":
        transcript = transcript_cache["data"]
    elif not meta["has_audio"]:
        transcript = {"status": "no_audio", "segments": [], "error": None}
    else:
        transcript = transcriber.run(video, CACHE / "audio" / f"{cid}.wav", OUTPUT / "transcripts" / f"{video.stem}.txt")
        atomic_json(transcript_path, {"source": source, "status": transcript["status"], "data": transcript})

    timeline = make_timeline(file_info, scenes, ocr_data, transcript)
    result = {
        "pipeline_version": PIPELINE_VERSION,
        "source": source,
        "status": "success",
        "error": None,
        "file": file_info,
        "quality": {
            "ffprobe": "success",
            "duration": file_info["duration"],
            "scenes": "success",
            "scene_count": len(scenes),
            "frames": "success",
            "keyframe_count": len(frames),
            "ocr": ocr_data.get("status"),
            "asr": transcript.get("status"),
            "event_partition": "success",
            "person_identification": "success" if timeline["persons"] else "无可靠实名人物",
            "needs_review": timeline["needs_review"],
            "notes": "",
        },
        "scenes": scenes,
        "keyframes": frames,
        "ocr": ocr_data,
        "transcript": transcript,
        **timeline,
        "processing_seconds": round(time.time() - started, 2),
    }
    atomic_json(out_json, result)
    logging.info("SUCCESS %s events=%d frames=%d %.1fs", video.name, len(result["events"]), len(frames), result["processing_seconds"])
    return result


def _join(values: list[Any]) -> str:
    return "；".join(str(v) for v in values if v not in (None, "", []))


def write_excel(records: list[dict[str, Any]]) -> None:
    wb = Workbook()
    total = wb.active
    total.title = "重大事件总时间轴"
    video_sheets = {
        "#250928-WDCC总结片（中大英小）": wb.create_sheet("#250928中大英小"),
        "#250928-WDCC总结片clena": wb.create_sheet("#250928clena"),
        "#250929-WDCC总结片（中小英大）": wb.create_sheet("#250929中小英大"),
    }
    people = wb.create_sheet("重要人物索引")
    quality = wb.create_sheet("处理质量")

    event_headers = ["序号", "视频名称", "事件编号", "开始时间", "结束时间", "持续时间", "事件类型", "事件标题",
                     "重要人物", "人物身份", "活动/机构", "地点", "事件描述", "重要语音摘要", "重要画面文字",
                     "检索关键词", "识别依据", "置信度", "是否需要人工复核"]
    people_headers = ["人物姓名/描述", "人物身份", "视频名称", "相关事件", "首次出现时间", "主要出现时间段", "识别依据", "置信度"]
    quality_headers = ["视频名称", "完整路径", "ffprobe状态", "视频时长", "场景检测状态", "关键帧数量", "OCR状态", "ASR状态",
                       "事件分区状态", "人物识别状态", "是否存在需要人工复核的事件", "异常说明"]

    for ws in [total, *video_sheets.values()]:
        ws.append(event_headers)
    people.append(people_headers)
    quality.append(quality_headers)

    serial = 1
    for record in records:
        f = record["file"]
        stem = Path(f["filename"]).stem
        target_ws = video_sheets[stem]
        for event in sorted(record["events"], key=lambda e: e["start_seconds"]):
            row = [
                serial, f["filename"], event["event_id"], event["start"], event["end"], event["duration"], event["type"],
                event["title"], _join(event["persons"]), _join(event["roles"]), event["activity_or_org"], event["location"],
                event["description"], event["speech_summary"], event["visual_text"], _join(event["keywords"]),
                _join(event["evidence"]), event["confidence"], "是" if event["needs_review"] else "否",
            ]
            total.append(row)
            target_ws.append(row)
            serial += 1
        for person in record.get("persons", []):
            people.append([person.get("name"), person.get("role"), f["filename"], _join(person.get("related_events", [])),
                           person.get("first_seen"), _join(person.get("time_ranges", [])), _join(person.get("evidence", [])),
                           person.get("confidence")])
        q = record["quality"]
        quality.append([f["filename"], f["path"], q["ffprobe"], sec_to_hms(q["duration"]), q["scenes"], q["keyframe_count"],
                        q["ocr"], q["asr"], q["event_partition"], q["person_identification"],
                        "是" if q["needs_review"] else "否", q.get("notes", "")])
        quality.cell(quality.max_row, 2).hyperlink = "file://" + quote(f["path"], safe="/")

    widths = {
        "重大事件总时间轴": [7, 28, 10, 12, 12, 12, 16, 32, 26, 24, 26, 18, 52, 48, 52, 42, 56, 10, 16],
        "#250928中大英小": [7, 28, 10, 12, 12, 12, 16, 32, 26, 24, 26, 18, 52, 48, 52, 42, 56, 10, 16],
        "#250928clena": [7, 28, 10, 12, 12, 12, 16, 32, 26, 24, 26, 18, 52, 48, 52, 42, 56, 10, 16],
        "#250929中小英大": [7, 28, 10, 12, 12, 12, 16, 32, 26, 24, 26, 18, 52, 48, 52, 42, 56, 10, 16],
        "重要人物索引": [22, 24, 28, 42, 14, 26, 58, 10],
        "处理质量": [28, 64, 14, 12, 16, 12, 14, 14, 16, 18, 22, 45],
    }
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{ws.cell(ws.max_row, ws.max_column).column_letter}{ws.max_row}"
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for idx, width in enumerate(widths[ws.title], 1):
            ws.column_dimensions[ws.cell(1, idx).column_letter].width = width
    EXCEL.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=EXCEL.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        check = load_workbook(tmp, read_only=False)
        expected = ["重大事件总时间轴", "#250928中大英小", "#250928clena", "#250929中小英大", "重要人物索引", "处理质量"]
        assert check.sheetnames == expected
        assert check["重大事件总时间轴"].max_row > 1
        check.close()
        os.replace(tmp, EXCEL)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def write_summary(records: list[dict[str, Any]], started: float) -> None:
    summary = {
        "root": str(ROOT),
        "videos": [
            {
                "filename": r["file"]["filename"],
                "path": r["file"]["path"],
                "duration": sec_to_hms(r["file"]["duration"]),
                "event_count": len(r["events"]),
                "review_event_count": sum(1 for e in r["events"] if e["needs_review"]),
                "person_count": len(r.get("persons", [])),
                "keyframe_count": r["quality"]["keyframe_count"],
                "ocr_status": r["quality"]["ocr"],
                "asr_status": r["quality"]["asr"],
            }
            for r in records
        ],
        "total_events": sum(len(r["events"]) for r in records),
        "total_persons": sum(len(r.get("persons", [])) for r in records),
        "excel_path": str(EXCEL),
        "json_dir": str(OUTPUT),
        "cache_dir": str(CACHE),
        "log_path": str(LOG),
        "elapsed_seconds": round(time.time() - started, 2),
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
    }
    atomic_json(OUTPUT / "run_summary.json", summary)


def main() -> None:
    setup_logging()
    started = time.time()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    paths = find_targets()
    logging.info("Targets: %s", [str(p) for p in paths])
    ocr_runner = OCRRunner()
    transcriber = Transcriber()
    records = [process_video(p, ocr_runner, transcriber) for p in paths]
    write_excel(records)
    write_summary(records, started)
    logging.info("DONE %.1fs Excel=%s", time.time() - started, EXCEL)


if __name__ == "__main__":
    main()
