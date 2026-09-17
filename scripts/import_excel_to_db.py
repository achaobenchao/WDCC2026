#!/usr/bin/env python3
"""Import the manually revised beta workbook into the local SQLite index.

This program reads the workbook only.  It never saves or modifies the Excel
file, and it performs no video analysis.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from database.db import DEFAULT_DB_PATH, initialize, normalize_text, parse_range, parse_timestamp  # noqa: E402


DEFAULT_EXCEL = Path("/Volumes/公用文件夹/WDCC2026/视频文件检索目录beta版/视频文件检索总表_beta汇总版.xlsx")


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def rows(ws):
    iterator = ws.iter_rows(values_only=True)
    headers = [text(v) for v in next(iterator, ())]
    for values in iterator:
        row = {headers[index]: text(values[index]) if index < len(values) else "" for index in range(len(headers))}
        if any(row.values()):
            yield row


def source_key(row: dict[str, str]) -> str:
    full_path = row.get("完整路径", "")
    if full_path:
        return f"path:{full_path}"
    return "fallback:" + "\x1f".join(row.get(key, "") for key in ("相对路径", "文件名", "视频时长"))


def find_video_id(conn, row: dict[str, str]) -> int | None:
    source_excel = row.get("来源表格", "")
    batch = row.get("所属批次 / 文件夹", "")
    full_path = row.get("完整路径", "")
    filename = row.get("视频文件", row.get("文件名", ""))
    if full_path:
        result = conn.execute(
            "SELECT id FROM videos WHERE source_excel=? AND source_batch=? AND full_path=? ORDER BY id LIMIT 1",
            (source_excel, batch, full_path),
        ).fetchone()
        if result:
            return result["id"]
    result = conn.execute(
        "SELECT id FROM videos WHERE source_excel=? AND source_batch=? AND filename=? ORDER BY id LIMIT 1",
        (source_excel, batch, filename),
    ).fetchone()
    return result["id"] if result else None


def split_keywords(value: str):
    for item in re.split(r"[；;、,，\n\r]+", value or ""):
        item = item.strip()
        if item:
            yield item


def upsert_video(conn, row: dict[str, str]) -> int:
    values = {
        "source_excel": row.get("来源表格", ""), "source_sheet": row.get("来源Sheet", ""),
        "source_batch": row.get("所属批次 / 文件夹", ""), "source_key": source_key(row),
        "filename": row.get("文件名", ""), "relative_path": row.get("相对路径", ""),
        "full_path": row.get("完整路径", ""), "folder": row.get("所在文件夹", ""),
        "file_size": row.get("文件大小", ""), "duration": row.get("视频时长", ""),
        "resolution": row.get("分辨率", ""), "fps": row.get("帧率", ""), "summary": row.get("视频内容摘要", ""),
        "activity_name": row.get("活动/会议名称", ""), "organization": row.get("重要机构", ""),
        "location": row.get("重要地点", ""), "important_date": row.get("重要日期", ""),
        "important_text": row.get("画面关键文字", ""), "speech_keywords": row.get("语音关键词", ""),
        "search_keywords": row.get("检索关键词", ""), "overall_confidence": row.get("总体置信度", ""),
        "review_required": row.get("是否需要人工复核", ""), "processing_status": row.get("处理状态", ""), "notes": row.get("备注", ""),
    }
    columns = list(values)
    update_columns = [column for column in columns if column not in {"source_excel", "source_sheet", "source_batch", "source_key"}]
    conn.execute(
        f"INSERT INTO videos ({','.join(columns)}) VALUES ({','.join('?' for _ in columns)}) "
        f"ON CONFLICT(source_excel,source_sheet,source_batch,source_key) DO UPDATE SET "
        + ",".join(f"{column}=excluded.{column}" for column in update_columns)
        + ",updated_at=CURRENT_TIMESTAMP",
        [values[column] for column in columns],
    )
    return conn.execute(
        "SELECT id FROM videos WHERE source_excel=? AND source_sheet=? AND source_batch=? AND source_key=?",
        (values["source_excel"], values["source_sheet"], values["source_batch"], values["source_key"]),
    ).fetchone()["id"]


def import_workbook(excel_path: Path, db_path: Path) -> dict[str, int]:
    conn = initialize(db_path)
    workbook = load_workbook(excel_path, read_only=True, data_only=False)
    required = {"素材总览", "人物索引", "事件索引"}
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"缺少必要工作表：{'、'.join(sorted(missing))}")
    counts = {"videos": 0, "persons": 0, "video_persons": 0, "events": 0, "keywords": 0, "sources": 0, "unmatched": 0}
    try:
        for row in rows(workbook["素材总览"]):
            video_id = upsert_video(conn, row)
            counts["videos"] += 1
            for field, kind in (("检索关键词", "search"), ("语音关键词", "speech"), ("画面关键文字", "ocr")):
                for keyword in split_keywords(row.get(field, "")):
                    conn.execute("INSERT OR IGNORE INTO keywords(video_id,keyword,keyword_type) VALUES (?,?,?)", (video_id, keyword, kind))
                    counts["keywords"] += 1
        for row in rows(workbook["人物索引"]):
            video_id = find_video_id(conn, row)
            if not video_id:
                counts["unmatched"] += 1
                continue
            name = row.get("人物姓名/描述", "")
            if not name:
                continue
            role = row.get("人物身份", "")
            normalized = normalize_text(name)
            conn.execute(
                "INSERT INTO persons(name,role,normalized_name) VALUES (?,?,?) "
                "ON CONFLICT(normalized_name) DO UPDATE SET role=CASE WHEN excluded.role<>'' THEN excluded.role ELSE persons.role END",
                (name, role, normalized),
            )
            person_id = conn.execute("SELECT id FROM persons WHERE normalized_name=?", (normalized,)).fetchone()["id"]
            conn.execute(
                "INSERT OR IGNORE INTO video_persons(video_id,person_id,first_seen,first_seen_seconds,time_ranges,related_events,evidence,confidence,source_excel,source_sheet) VALUES (?,?,?,?,?,?,?,?,?,?)",
                (video_id, person_id, row.get("首次出现时间", ""), parse_timestamp(row.get("首次出现时间", "")), row.get("主要出现时间", ""), row.get("相关事件", ""), row.get("识别依据", ""), row.get("置信度", ""), row.get("来源表格", ""), row.get("来源Sheet", "")),
            )
            counts["persons"] += 1; counts["video_persons"] += 1
        for row in rows(workbook["事件索引"]):
            video_id = find_video_id(conn, row)
            if not video_id:
                counts["unmatched"] += 1
                continue
            start_text = row.get("开始时间", "")
            end_text = row.get("结束时间", "")
            start_seconds, range_end_seconds = parse_range(start_text)
            # A normal start-time cell is a single timestamp.  In that case the
            # event's end comes from the separate end-time column.  If the start
            # cell itself holds a range, preserve its range endpoint instead.
            if any(separator in start_text for separator in ("-", "–", "—")):
                end_seconds = range_end_seconds
            else:
                _, end_seconds = parse_range(end_text)
            video = conn.execute("SELECT organization,location,review_required FROM videos WHERE id=?", (video_id,)).fetchone()
            conn.execute(
                "INSERT INTO events(video_id,event_number,event_type,event_name,start_time,end_time,start_seconds,end_seconds,description,persons_text,organization,location,evidence,confidence,review_required,source_excel,source_sheet) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) "
                "ON CONFLICT(video_id,event_number,event_type,event_name,start_time,end_time,source_excel,source_sheet) DO UPDATE SET description=excluded.description,persons_text=excluded.persons_text,evidence=excluded.evidence,confidence=excluded.confidence,updated_at=CURRENT_TIMESTAMP",
                (video_id, row.get("事件编号", ""), row.get("事件类型", ""), row.get("事件名称", ""), row.get("开始时间", ""), row.get("结束时间", ""), start_seconds, end_seconds, row.get("事件描述", ""), row.get("涉及人物", ""), video["organization"], video["location"], row.get("识别依据", ""), row.get("置信度", ""), video["review_required"], row.get("来源表格", ""), row.get("来源Sheet", "")),
            )
            counts["events"] += 1
        if "来源说明" in workbook.sheetnames:
            for row in rows(workbook["来源说明"]):
                source_path = row.get("完整路径", "")
                if not source_path:
                    continue
                conn.execute(
                    "INSERT INTO import_sources(source_filename,source_path,import_time,rows_read,rows_imported,rows_skipped,duplicate_count,notes) VALUES (?,?,CURRENT_TIMESTAMP,?,?,?,?,?) "
                    "ON CONFLICT(source_path) DO UPDATE SET import_time=CURRENT_TIMESTAMP,rows_read=excluded.rows_read,rows_imported=excluded.rows_imported,rows_skipped=excluded.rows_skipped,duplicate_count=excluded.duplicate_count,notes=excluded.notes",
                    (row.get("源Excel文件名", ""), source_path, row.get("读取行数", ""), row.get("成功导入行数", ""), row.get("跳过行数", ""), row.get("疑似重复数", ""), row.get("异常说明", "")),
                )
                counts["sources"] += 1
        conn.commit()
    finally:
        conn.close()
    return counts


def main() -> None:
    parser = argparse.ArgumentParser(description="导入人工修订的 Excel 总表至 SQLite（只读 Excel）。")
    parser.add_argument("excel", nargs="?", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH)
    args = parser.parse_args()
    if not args.excel.is_file():
        parser.error(f"找不到 Excel：{args.excel}")
    result = import_workbook(args.excel, args.db)
    print("导入完成：" + "；".join(f"{key}={value}" for key, value in result.items()))


if __name__ == "__main__":
    main()
