#!/usr/bin/env python3
"""普通视频素材分析：只处理“视频总结片/视频总结片”根目录下 7 个日期文件夹。

Excluded by design:
- root-level three long summary videos already indexed separately
- non-date sibling folders
- any NAS writes
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from video_indexer.analyze import fuse
from video_indexer.cache import atomic_json, cache_id, fingerprint, load_json, reusable
from video_indexer.config import VIDEO_EXTENSIONS
from video_indexer.frames import extract
from video_indexer.main import OCR_VERSION, PIPELINE_VERSION
from video_indexer.ocr import OCRRunner
from video_indexer.transcribe import Transcriber
from video_indexer.video_info import probe


ROOT = Path("/Volumes/公用文件夹/WDCC2026/视频总结片/视频总结片")
DATE_FOLDERS = ["9.19", "9.21", "9.24", "9.25", "9.26", "9.27", "9.28"]
RUN_NAME = "视频总结片_7文件夹"
EXCEL_NAME = "视频素材检索结果_视频总结片_7文件夹.xlsx"
STATUS_SUCCESS = {"SUCCESS", "success"}
STATUS_FAILED = {"FAILED", "FAILED_NATIVE_CRASH", "failed"}


def setup_logging(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[logging.FileHandler(path, encoding="utf-8"), logging.StreamHandler()],
        force=True,
    )


def scan_selected() -> list[dict]:
    videos = []
    for folder in DATE_FOLDERS:
        folder_path = ROOT / folder
        if not folder_path.is_dir():
            logging.warning("missing folder: %s", folder_path)
            continue
        for path in folder_path.rglob("*"):
            if path.is_file() and path.suffix.lower() in VIDEO_EXTENSIONS:
                resolved = path.resolve()
                stat = resolved.stat()
                videos.append({
                    "filename": resolved.name,
                    "path": str(resolved),
                    "relative_path": str(resolved.relative_to(ROOT)),
                    "folder": str(resolved.parent.relative_to(ROOT)),
                    "top_folder": folder,
                    "size": stat.st_size,
                    "modified_time": stat.st_mtime,
                    "mtime_ns": stat.st_mtime_ns,
                    "status": "pending",
                })
    videos.sort(key=lambda item: item["relative_path"].casefold())
    return videos


def fmt_seconds(seconds: float | int | None) -> str:
    seconds = max(0, int(round(seconds or 0)))
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def join(values) -> str:
    return "；".join(str(v) for v in values if v not in (None, "", []))


def flush_logs() -> None:
    for handler in logging.getLogger().handlers:
        try:
            handler.flush()
        except Exception:
            pass


def stage_start(base: dict, result_path: Path, stage: str) -> None:
    base["current_stage"] = stage
    base.setdefault("stage_state", {})[f"{stage}_done"] = False
    logging.info("STAGE_START %s %s", stage, base.get("file", {}).get("relative_path"))
    flush_logs()
    atomic_json(result_path, base)


def stage_done(base: dict, result_path: Path, stage: str) -> None:
    base.setdefault("stage_state", {})[f"{stage}_done"] = True
    base["last_successful_stage"] = stage
    logging.info("STAGE_DONE %s %s", stage, base.get("file", {}).get("relative_path"))
    flush_logs()
    atomic_json(result_path, base)


def is_success(record: dict | None) -> bool:
    return bool(record and record.get("status") in STATUS_SUCCESS and record.get("source"))


def process_one(item: dict, cache: Path, output: Path, ocr_runner: OCRRunner, transcriber: Transcriber, resume: bool, attempt: int = 0) -> dict:
    path = Path(item["path"])
    cid = cache_id(path)
    source = fingerprint(path)
    result_path = output / "index" / f"{cid}.json"
    existing = load_json(result_path)
    if resume and reusable(existing, source) and existing.get("pipeline_version") == PIPELINE_VERSION and existing.get("status") in STATUS_SUCCESS:
        return existing

    started = time.time()
    quality = {
        "ffprobe": "pending", "scenes": "pending", "frames": "pending", "ocr": "pending",
        "transcription": "pending", "persons": "pending", "events": "pending", "errors": "",
    }
    base = {
        "pipeline_version": PIPELINE_VERSION,
        "source": source,
        "file": item,
        "persons": [],
        "events": [],
        "organizations": [],
        "locations": [],
        "activities": [],
        "ocr_keywords": [],
        "speech_keywords": [],
        "search_keywords": [],
        "summary": "",
        "confidence": "低",
        "needs_review": True,
        "quality": quality,
        "stage_state": {
            "metadata_done": False, "scene_done": False, "frames_done": False,
            "ocr_done": False, "transcription_done": False, "analysis_done": False, "json_done": False,
        },
        "current_stage": "init",
        "last_successful_stage": existing.get("last_successful_stage") if existing else None,
        "attempt": attempt,
        "status": "PROCESSING",
        "error": None,
    }
    logging.info("VIDEO_START %s attempt=%d", item["relative_path"], attempt)
    flush_logs()
    atomic_json(result_path, base)
    errors = []
    try:
        stage_start(base, result_path, "metadata")
        meta_path = cache / "metadata" / f"{cid}.json"
        meta_cache = load_json(meta_path)
        if resume and reusable(meta_cache, source):
            meta = meta_cache["data"]
        else:
            meta = probe(path)
            atomic_json(meta_path, {"source": source, "status": "success", "data": meta})
        base["file"].update(meta)
        quality["ffprobe"] = "success"
        stage_done(base, result_path, "metadata")
    except Exception as exc:
        quality["ffprobe"] = "failed"
        errors.append(f"ffprobe: {exc}")
        base.update(status="FAILED", error=" | ".join(errors))
        quality["errors"] = base["error"]
        atomic_json(result_path, base)
        return base

    try:
        stage_start(base, result_path, "scene")
        scene_path = cache / "scenes" / f"{cid}.json"
        scene_cache = load_json(scene_path)
        if resume and reusable(scene_cache, source) and scene_cache.get("strategy") == "timeline_sampling":
            candidates = scene_cache["candidates"]
        else:
            duration = float(meta["duration"])
            scenes, candidates = timeline_sample_candidates(duration)
            atomic_json(scene_path, {"source": source, "status": "success", "strategy": "timeline_sampling", "scenes": scenes, "candidates": candidates})
        candidates = cap_keyframe_candidates(candidates, meta["duration"])
        quality["scenes"] = "timeline_sampling"
        stage_done(base, result_path, "scene")
        stage_start(base, result_path, "frames")
        frames = extract(path, candidates, cache / "frames" / cid)
        quality["frames"] = "success" if frames else "failed"
        if not frames:
            errors.append("未提取到可读关键帧")
        else:
            stage_done(base, result_path, "frames")
    except Exception as exc:
        frames = []
        quality["scenes"] = "failed"
        quality["frames"] = "failed"
        errors.append(f"场景/关键帧: {exc}")

    stage_start(base, result_path, "ocr")
    ocr_path = cache / "ocr" / f"{cid}.json"
    ocr_cache = load_json(ocr_path)
    if resume and reusable(ocr_cache, source) and ocr_cache.get("ocr_version") == OCR_VERSION:
        ocr_data = ocr_cache["data"]
    elif frames:
        ocr_data = ocr_runner.run(frames)
        atomic_json(ocr_path, {"ocr_version": OCR_VERSION, "source": source, "status": ocr_data["status"], "data": ocr_data})
    else:
        ocr_data = {"status": "skipped", "error": "无关键帧", "frames": []}
    quality["ocr"] = ocr_data["status"]
    if ocr_data.get("error"):
        errors.append("OCR: " + ocr_data["error"])
    if ocr_data["status"] == "success":
        stage_done(base, result_path, "ocr")

    stage_start(base, result_path, "transcription")
    transcript_path = cache / "transcripts" / f"{cid}.json"
    transcript_cache = load_json(transcript_path)
    if not meta["has_audio"]:
        transcript = {"status": "no_audio", "segments": [], "error": None}
    elif resume and reusable(transcript_cache, source):
        transcript = transcript_cache["data"]
    else:
        transcript = transcriber.run(path, cache / "audio" / f"{cid}.wav", output / "transcripts" / f"{cid}.txt")
        atomic_json(transcript_path, {"source": source, "status": transcript["status"], "data": transcript})
    quality["transcription"] = transcript["status"]
    if transcript.get("error"):
        errors.append("ASR: " + transcript["error"])
    if transcript["status"] in {"success", "no_audio"}:
        stage_done(base, result_path, "transcription")

    stage_start(base, result_path, "analysis")
    analysis = fuse(base["file"], ocr_data, transcript)
    base.update(analysis)
    quality["persons"] = "success" if base["persons"] else "无重要人物"
    quality["events"] = "success" if base["events"] else "无明确重要事件"
    stage_done(base, result_path, "analysis")
    quality["errors"] = " | ".join(errors)
    required_ok = (
        quality["ffprobe"] == "success"
        and quality["scenes"] in {"success", "timeline_sampling"}
        and quality["frames"] == "success"
        and quality["ocr"] == "success"
        and quality["transcription"] in {"success", "no_audio"}
    )
    base["status"] = "SUCCESS" if required_ok else "FAILED"
    base["error"] = quality["errors"] or (None if required_ok else "必要阶段未全部完成")
    base["processing_seconds"] = round(time.time() - started, 2)
    base["current_stage"] = "json"
    base.setdefault("stage_state", {})["json_done"] = False
    atomic_json(result_path, base)
    verified = load_json(result_path)
    if verified and verified.get("status") == base["status"] and verified.get("source") == source:
        base.setdefault("stage_state", {})["json_done"] = True
        base["current_stage"] = "complete" if required_ok else "failed"
        atomic_json(result_path, base)
    logging.info("VIDEO_DONE %s status=%s seconds=%.1f", item["relative_path"], base["status"], base["processing_seconds"])
    flush_logs()
    return base


def cap_keyframe_candidates(candidates: list[dict], duration: float) -> list[dict]:
    """Limit NAS random seeks for ordinary short clips while preserving timeline coverage."""
    if not candidates:
        return candidates
    if duration <= 15:
        limit = 1
    elif duration <= 30:
        limit = 2
    elif duration <= 60:
        limit = 4
    elif duration <= 120:
        limit = 5
    else:
        limit = 8
    if len(candidates) <= limit:
        return candidates
    ordered = sorted(candidates, key=lambda item: float(item["timestamp"]))
    indexes = sorted({round(i * (len(ordered) - 1) / (limit - 1)) for i in range(limit)})
    return [ordered[i] for i in indexes]


def timeline_sample_candidates(duration: float) -> tuple[list[dict], list[dict]]:
    """Fast representative-frame plan for large batches of ordinary short clips."""
    duration = max(0.0, float(duration or 0))
    if duration <= 0:
        return [], []
    if duration <= 15:
        stamps = [duration / 2]
    elif duration <= 30:
        stamps = [duration * 0.35, duration * 0.75]
    elif duration <= 60:
        stamps = [duration * 0.25, duration * 0.5, duration * 0.8]
    elif duration <= 120:
        stamps = [duration * 0.15, duration * 0.35, duration * 0.6, duration * 0.85]
    else:
        stamps = [duration * x for x in (0.08, 0.2, 0.35, 0.5, 0.65, 0.8, 0.92)]
    candidates = [
        {"timestamp": max(0.0, min(duration - 0.1, stamp)), "scene": i, "reason": "timeline_sample"}
        for i, stamp in enumerate(stamps, 1)
    ]
    return [{"scene": 1, "start": 0.0, "end": duration}], candidates


def export_excel(records: list[dict], excel_path: Path) -> None:
    wb = Workbook()
    overview = wb.active
    overview.title = "素材总览"
    people = wb.create_sheet("人物索引")
    events = wb.create_sheet("事件索引")
    quality = wb.create_sheet("处理质量")

    overview_headers = ["序号", "所属子文件夹", "文件名", "相对路径", "完整路径", "所在文件夹", "文件大小", "视频时长", "分辨率", "帧率",
                        "重要人物", "人物身份", "人物出现时间", "主要事件", "事件时间", "视频内容摘要", "活动/会议名称", "重要机构",
                        "重要地点", "画面关键文字", "语音关键词", "检索关键词", "总体置信度", "是否需要人工复核", "处理状态", "备注"]
    people_headers = ["人物姓名", "人物身份", "视频文件", "所属子文件夹", "相对路径", "完整路径", "首次出现时间", "主要出现时间", "相关事件", "识别依据", "置信度"]
    events_headers = ["视频文件", "所属子文件夹", "相对路径", "完整路径", "事件编号", "事件类型", "事件名称", "涉及人物", "开始时间", "结束时间", "事件描述", "识别依据", "置信度"]
    quality_headers = ["所属子文件夹", "文件名", "ffprobe状态", "场景检测状态", "关键帧状态", "OCR状态", "语音识别状态", "人物分析状态", "事件分析状态", "整体状态", "是否需要人工复核", "异常说明"]

    for sheet, headers in [(overview, overview_headers), (people, people_headers), (events, events_headers), (quality, quality_headers)]:
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).column_letter}1"
        sheet.sheet_view.showGridLines = False

    for i, r in enumerate(records, 1):
        f, q = r.get("file", {}), r.get("quality", {})
        ps, es = r.get("persons", []), r.get("events", [])
        notes = join([r.get("error"), q.get("errors")])
        overview.append([i, f.get("top_folder") or str(f.get("relative_path", "")).split("/", 1)[0], f.get("filename"), f.get("relative_path"), f.get("path"),
                         f.get("folder"), f.get("size"), (f.get("duration") or 0) / 86400, f.get("resolution"), f.get("fps"),
                         join([p.get("name") for p in ps]) or "无重要人物", join([p.get("role") for p in ps]),
                         join([join(p.get("time_ranges", [])) for p in ps]), join([e.get("type") for e in es]) or "无明确重要事件",
                         join([f'{e.get("start", "")}-{e.get("end", "")}' for e in es]), r.get("summary"), join(r.get("activities", [])),
                         join(r.get("organizations", [])), join(r.get("locations", [])), join(r.get("ocr_keywords", [])),
                         join(r.get("speech_keywords", [])), join(r.get("search_keywords", [])), r.get("confidence"),
                         "是" if r.get("needs_review") else "否", r.get("status"), notes])
        overview.cell(overview.max_row, 5).hyperlink = "file://" + quote(f.get("path", ""), safe="/")
        for p in ps:
            people.append([p.get("name"), p.get("role"), f.get("filename"), f.get("top_folder"), f.get("relative_path"), f.get("path"), p.get("first_seen"),
                           join(p.get("time_ranges", [])), join(p.get("related_events", [])), join(p.get("evidence", [])), p.get("confidence")])
            people.cell(people.max_row, 6).hyperlink = "file://" + quote(f.get("path", ""), safe="/")
        for n, e in enumerate(es, 1):
            events.append([f.get("filename"), f.get("top_folder"), f.get("relative_path"), f.get("path"), n, e.get("type"), e.get("name"),
                           join(e.get("persons", [])), e.get("start"), e.get("end"), e.get("description"), join(e.get("evidence", [])), e.get("confidence")])
            events.cell(events.max_row, 4).hyperlink = "file://" + quote(f.get("path", ""), safe="/")
        quality.append([f.get("top_folder"), f.get("filename"), q.get("ffprobe"), q.get("scenes"), q.get("frames"), q.get("ocr"), q.get("transcription"),
                        q.get("persons"), q.get("events"), r.get("status"), "是" if r.get("needs_review") else "否", notes])

    widths = {
        "素材总览": [7, 16, 28, 45, 60, 32, 14, 12, 13, 9, 24, 24, 24, 24, 24, 42, 22, 24, 20, 45, 40, 40, 12, 16, 12, 35],
        "人物索引": [22, 25, 28, 16, 45, 60, 14, 25, 25, 45, 10],
        "事件索引": [28, 16, 45, 60, 10, 16, 26, 24, 14, 14, 50, 50, 10],
        "处理质量": [16, 28, 14, 16, 14, 14, 16, 16, 16, 14, 18, 45],
    }
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col, width in enumerate(widths[sheet.title], 1):
            sheet.column_dimensions[sheet.cell(1, col).column_letter].width = width
        sheet.auto_filter.ref = f"A1:{sheet.cell(sheet.max_row, sheet.max_column).column_letter}{sheet.max_row}"
    for cell in overview["H"][1:]:
        cell.number_format = "[h]:mm:ss"

    excel_path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=excel_path.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        check = load_workbook(tmp)
        assert check.sheetnames == ["素材总览", "人物索引", "事件索引", "处理质量"]
        assert check["素材总览"].max_row == len(records) + 1
        check.close()
        os.replace(tmp, excel_path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def summarize(records: list[dict], videos: list[dict], started: float, output: Path, cache: Path, log_path: Path, excel_path: Path) -> dict:
    folder_counts = {}
    folder_failures = {}
    for r in records:
        top = r.get("file", {}).get("top_folder") or str(r.get("file", {}).get("relative_path", "")).split("/", 1)[0]
        folder_counts[top] = folder_counts.get(top, 0) + 1
        if r.get("status") in STATUS_FAILED:
            folder_failures[top] = folder_failures.get(top, 0) + 1
    return {
        "root": str(ROOT),
        "folders": DATE_FOLDERS,
        "video_total": len(videos),
        "records": len(records),
        "success": sum(r.get("status") in STATUS_SUCCESS for r in records),
        "failed": sum(r.get("status") in STATUS_FAILED for r in records),
        "needs_review": sum(bool(r.get("needs_review")) for r in records),
        "persons_effective": sum(bool(r.get("persons")) for r in records),
        "events_effective": sum(bool(r.get("events")) for r in records),
        "folder_counts": folder_counts,
        "folder_failures": folder_failures,
        "elapsed_seconds": round(time.time() - started, 2),
        "excel_path": str(excel_path),
        "json_dir": str(output / "index"),
        "cache_dir": str(cache),
        "log_path": str(log_path),
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
    }


def status_counts(videos: list[dict], output: Path) -> dict:
    counts = {"total": len(videos), "success": 0, "pending": 0, "processing": 0, "failed": 0}
    for item in videos:
        record = load_json(output / "index" / f'{cache_id(Path(item["path"]))}.json')
        if not record:
            counts["pending"] += 1
        elif record.get("status") in STATUS_SUCCESS:
            counts["success"] += 1
        elif record.get("status") in STATUS_FAILED:
            counts["failed"] += 1
        elif record.get("status") == "PROCESSING":
            counts["processing"] += 1
        else:
            counts["pending"] += 1
    return counts


def mark_native_crash(item: dict, output: Path, attempt: int, returncode: int, final: bool) -> None:
    path = Path(item["path"])
    cid = cache_id(path)
    source = fingerprint(path)
    result_path = output / "index" / f"{cid}.json"
    record = load_json(result_path) or {
        "pipeline_version": PIPELINE_VERSION,
        "source": source,
        "file": item,
        "quality": {},
        "stage_state": {},
        "status": "PROCESSING",
    }
    record["worker_crash_count"] = int(record.get("worker_crash_count") or 0) + 1
    record["native_exit_code"] = returncode
    record["failure_type"] = "FAILED_NATIVE_CRASH"
    record["error"] = f"worker native/abnormal exit code {returncode}; attempt={attempt}; last_stage={record.get('current_stage')}; last_successful_stage={record.get('last_successful_stage')}"
    if final:
        record["status"] = "FAILED"
        record.setdefault("quality", {})["errors"] = record["error"]
    else:
        record["status"] = "PROCESSING"
    atomic_json(result_path, record)


def run_worker_for_item(item_index: int, item: dict, args, project: Path, log_path: Path) -> None:
    result_path = project / "output" / RUN_NAME / "index" / f'{cache_id(Path(item["path"]))}.json'
    existing = load_json(result_path)
    if existing and existing.get("source") == fingerprint(Path(item["path"])) and existing.get("status") in STATUS_SUCCESS:
        logging.info("SKIP_SUCCESS %s", item["relative_path"])
        return
    if existing and existing.get("status") in STATUS_FAILED:
        logging.info("SKIP_FAILED %s status=%s", item["relative_path"], existing.get("failure_type") or existing.get("status"))
        return
    for attempt in range(2):
        logging.info("WORKER_START index=%d attempt=%d %s", item_index, attempt, item["relative_path"])
        flush_logs()
        cmd = [
            sys.executable, str(Path(__file__).resolve()),
            "--worker-index", str(item_index),
            "--worker-attempt", str(attempt),
            "--folders", ",".join(DATE_FOLDERS),
            "--run-name", RUN_NAME,
            "--excel-name", EXCEL_NAME,
        ]
        if not args.resume:
            cmd.append("--no-resume")
        result = subprocess.run(cmd, cwd=str(project))
        if result.returncode == 0:
            record = load_json(result_path)
            logging.info("WORKER_DONE index=%d status=%s %s", item_index, record.get("status") if record else "missing", item["relative_path"])
            flush_logs()
            return
        final = attempt >= 1
        logging.error("WORKER_CRASH index=%d attempt=%d returncode=%s %s", item_index, attempt, result.returncode, item["relative_path"])
        mark_native_crash(item, project / "output" / RUN_NAME, attempt, result.returncode, final)
        flush_logs()


def main() -> None:
    global DATE_FOLDERS, RUN_NAME, EXCEL_NAME
    parser = argparse.ArgumentParser()
    parser.add_argument("--shard-count", type=int, default=1)
    parser.add_argument("--shard-index", type=int, default=0)
    parser.add_argument("--resume", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--export-only", action="store_true")
    parser.add_argument("--limit", type=int)
    parser.add_argument("--worker-index", type=int)
    parser.add_argument("--worker-attempt", type=int, default=0)
    parser.add_argument("--folders", default=",".join(DATE_FOLDERS), help="Comma-separated top-level folders under ROOT, e.g. 9.27")
    parser.add_argument("--run-name", default=RUN_NAME, help="Local cache/output/log namespace")
    parser.add_argument("--excel-name", default=EXCEL_NAME, help="Excel filename under output/")
    args = parser.parse_args()
    if args.shard_count < 1 or not 0 <= args.shard_index < args.shard_count:
        parser.error("invalid shard")
    DATE_FOLDERS = [part.strip() for part in args.folders.split(",") if part.strip()]
    if not DATE_FOLDERS:
        parser.error("--folders must include at least one folder")
    RUN_NAME = args.run_name
    EXCEL_NAME = args.excel_name

    project = Path.cwd()
    cache = project / "cache" / RUN_NAME
    output = project / "output" / RUN_NAME
    log_path = project / "logs" / RUN_NAME / "process.log"
    excel_path = project / "output" / EXCEL_NAME
    setup_logging(log_path)
    started = time.time()
    manifest_path = output / "manifest.json"
    manifest = load_json(manifest_path)
    if manifest and manifest.get("root") == str(ROOT) and manifest.get("folders") == DATE_FOLDERS and manifest.get("videos"):
        videos = manifest["videos"]
    else:
        videos = scan_selected()
    if args.limit:
        videos = videos[:args.limit]
    atomic_json(output / "manifest.json", {"root": str(ROOT), "folders": DATE_FOLDERS, "count": len(videos), "videos": videos})

    if args.worker_index is not None:
        if not 0 <= args.worker_index < len(videos):
            parser.error("--worker-index out of range")
        item = videos[args.worker_index]
        ocr_runner, transcriber = OCRRunner(), Transcriber()
        record = process_one(item, cache, output, ocr_runner, transcriber, args.resume, args.worker_attempt)
        raise SystemExit(0 if record.get("status") in STATUS_SUCCESS | STATUS_FAILED else 2)

    logging.info("selected folders: %s", DATE_FOLDERS)
    logging.info("selected videos: %d", len(videos))
    counts = status_counts(videos, output)
    logging.info("RESUME_STATUS total=%d success=%d pending=%d processing_to_recover=%d failed=%d",
                 counts["total"], counts["success"], counts["pending"], counts["processing"], counts["failed"])

    if not args.export_only:
        work = [(i, v) for i, v in enumerate(videos) if i % args.shard_count == args.shard_index]
        logging.info("shard %d/%d videos=%d", args.shard_index + 1, args.shard_count, len(work))
        for pos, (item_index, item) in enumerate(work, 1):
            logging.info("[%d/%d] Processing %s", pos, len(work), item["relative_path"])
            run_worker_for_item(item_index, item, args, project, log_path)

    records = []
    for item in videos:
        record = load_json(output / "index" / f'{cache_id(Path(item["path"]))}.json')
        if record:
            record.setdefault("file", {}).setdefault("top_folder", item["top_folder"])
            records.append(record)
    if not args.limit and (args.export_only or len(records) == len(videos)):
        export_excel(records, excel_path)
    summary = summarize(records, videos, started, output, cache, log_path, excel_path)
    atomic_json(output / "run_summary.json", summary)
    logging.info("summary: %s", json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
